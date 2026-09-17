"""
Exportacion a Excel.

REEMPLAZA a la automatizacion COM del formulario original
(CREATEOBJECT("Excel.Application")), que exigia tener Excel instalado en la
maquina donde corre el programa y dejaba procesos EXCEL.EXE colgados si algo
fallaba a mitad de camino.

openpyxl escribe el .xlsx directamente. No hace falta Excel para generarlo.

Mapa de columnas (el del formulario, mas Ubicacion):
A Codigo | B Cod. Prov. | C Descripcion | D Ubicacion | E Und. | F Param1 |
G Param2 | H Stock local | I Stock otra suc. | J Stock total | K A pedir |
L Costo | M Subtotal
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..calculo import MODO_ROTACION, Linea
from .comun import nombre_archivo

AZUL = "FF1F4E79"
GRIS = "FFF2F2F2"
VERDE = "FFE2EFDA"


def exportar_pedido(
    lineas: list[Linea], proveedor, parametros, modo: int
) -> Path:
    ruta = nombre_archivo(
        "Pedido", proveedor.nombre if proveedor else "SIN_PROVEEDOR", ".xlsx"
    )

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Pedido"

    es_rotacion = modo == MODO_ROTACION
    nombre_proveedor = proveedor.nombre.strip() if proveedor else "Sin proveedor"

    # --- encabezado -----------------------------------------------------
    hoja["A1"] = "PEDIDO DE REPOSICION"
    hoja["A1"].font = Font(size=14, bold=True, color="FFFFFFFF")
    hoja["A1"].fill = PatternFill("solid", fgColor=AZUL)
    hoja.merge_cells("A1:M1")
    hoja["A1"].alignment = Alignment(horizontal="left", vertical="center")
    hoja.row_dimensions[1].height = 24

    import datetime

    hoja["A2"] = (
        f"Proveedor: {nombre_proveedor}   |   "
        f"Sucursal: {parametros.sucursal}   |   "
        f"Fecha: {datetime.date.today().strftime('%d/%m/%Y')}   |   "
        + ("Modo rotacion" if es_rotacion else "Modo min/max")
    )
    hoja["A2"].font = Font(size=10, italic=True)
    hoja.merge_cells("A2:M2")

    # AR_UBICACI dice donde esta fisicamente la mercaderia. En una lista de
    # compra sirve para ir a mirar el estante antes de pedir, sobre todo en
    # los articulos que figuran en negativo.
    encabezados = [
        "Codigo", "Cod. Prov.", "Descripcion", "Ubicacion", "Und.",
        "Vta/mes" if es_rotacion else "Minimo",
        "Dias stk" if es_rotacion else "Maximo",
        "Stock local", "Stock otra suc.", "Stock total",
        "A pedir", "Costo", "Subtotal",
    ]

    fila_encabezado = 4
    for columna, titulo in enumerate(encabezados, start=1):
        celda = hoja.cell(row=fila_encabezado, column=columna, value=titulo)
        celda.font = Font(bold=True, color="FFFFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    borde = Side(style="thin", color="FFBFBFBF")
    marco = Border(left=borde, right=borde, top=borde, bottom=borde)

    fila = fila_encabezado + 1
    for linea in sorted(lineas, key=lambda l: l.descripcion):
        valores = [
            linea.codigo,
            # Texto a proposito: si el codigo del proveedor tiene ceros a la
            # izquierda, Excel se los come al interpretarlo como numero.
            linea.cod_prov_articulo,
            linea.descripcion,
            linea.ubicacion,
            linea.unidad,
            linea.vta_mes if es_rotacion else linea.minimo,
            linea.dias_stock if es_rotacion else linea.maximo,
            linea.stock_local,
            linea.stock_remoto,
            linea.stock_total,
            linea.cant_pedir,
            linea.costo,
            linea.subtotal,
        ]
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(row=fila, column=columna, value=valor)
            celda.border = marco
            if columna in (1, 2, 3):
                celda.alignment = Alignment(horizontal="left")
            elif columna in (4, 5):
                celda.alignment = Alignment(horizontal="center")
            else:
                celda.alignment = Alignment(horizontal="right")
            if columna in (6, 7, 8, 9, 10):
                celda.number_format = "#,##0.00"
            elif columna == 11:
                celda.number_format = "#,##0"
                celda.font = Font(bold=True)
                celda.fill = PatternFill("solid", fgColor=VERDE)
            elif columna in (12, 13):
                celda.number_format = '"$" #,##0.00'
        hoja.cell(row=fila, column=2).number_format = "@"
        fila += 1

    # --- totales ---------------------------------------------------------
    total = hoja.cell(row=fila, column=10, value="TOTALES")
    total.font = Font(bold=True)
    total.alignment = Alignment(horizontal="right")

    celda_unidades = hoja.cell(
        row=fila, column=11,
        value=f"=SUM(K{fila_encabezado + 1}:K{fila - 1})",
    )
    celda_unidades.font = Font(bold=True)
    celda_unidades.number_format = "#,##0"

    celda_monto = hoja.cell(
        row=fila, column=13,
        value=f"=SUM(M{fila_encabezado + 1}:M{fila - 1})",
    )
    celda_monto.font = Font(bold=True)
    celda_monto.number_format = '"$" #,##0.00'

    doble = Side(style="double", color="FF000000")
    for columna in range(1, 14):
        hoja.cell(row=fila, column=columna).border = Border(top=doble)
        hoja.cell(row=fila, column=columna).fill = PatternFill("solid", fgColor=GRIS)

    # --- formato general -------------------------------------------------
    anchos = [12, 14, 46, 12, 7, 11, 10, 13, 15, 12, 10, 14, 16]
    for columna, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(columna)].width = ancho

    hoja.freeze_panes = f"A{fila_encabezado + 1}"
    hoja.auto_filter.ref = f"A{fila_encabezado}:M{fila - 1}"
    hoja.print_title_rows = f"{fila_encabezado}:{fila_encabezado}"
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.fitToWidth = 1

    libro.save(ruta)
    return ruta
